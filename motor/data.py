"""Läser speldata ur kortdata/*.xlsx — Excel är källan (se OMSTART.md)."""
from pathlib import Path

from openpyxl import load_workbook

KORTDATA = Path(__file__).resolve().parent.parent / "kortdata"


def las_lek(filnamn):
    """Alla rader på fliken Kort som dict (rubrik -> värde)."""
    ws = load_workbook(KORTDATA / filnamn, read_only=True, data_only=True)["Kort"]
    rader = ws.iter_rows(values_only=True)
    rubriker = next(rader)
    return [dict(zip(rubriker, r)) for r in rader if any(v not in (None, "") for v in r)]


def tal(v, standard=0.0):
    """'+0,5' / '-1' / 2 / None -> float."""
    if v in (None, ""):
        return standard
    return float(str(v).replace(",", ".").replace("−", "-"))


class Kortdata:
    """Alla lekar som behövs för Skede 3."""

    def __init__(self):
        self.handelse = las_lek("F_händelsekort.xlsx")
        self.kvartal = las_lek("F_kvartalskort.xlsx")
        self.person = las_lek("F_personkort.xlsx")
        self.omvarld = las_lek("F_omvärldskort.xlsx")
        self.dd = las_lek("F_DD.xlsx")
        self.yieldkort = las_lek("F_yieldkort.xlsx")
        self.fc = las_lek("F_FC.xlsx")
        self.fs = las_lek("F_FS.xlsx")
        alla = las_lek("PU_projekt.xlsx")
        self.projekt = [p for p in alla if p["Typ"] != "BRF"]
        self.brf = [p for p in alla if p["Typ"] == "BRF"]
