"""Gemensam byggare för kortdata-Excel (en fil per korttyp i kortdata/).

Varje korttyp har ett eget skript verktyg/bygg_<korttyp>_excel.py som
  1. läser CSV:n som InDesign-mallen fylldes med (arv/spelet2_v1/...),
  2. läser text utläst ur tryckfilen (arv/tryckt_text/<korttyp>_tryckeri.txt),
  3. jämför dem och avbryter vid avvikelser i speldata,
  4. anropar bygg_arbetsbok() nedan.

Alla Excel-filer får samma flikar:
  Kort               en rad per unikt kort — tryckta fält först, sedan ej tryckt speldata
  Mallens fasta text text som står i InDesign-mallen men inte i CSV:n
  Produktion         bildsökvägar, färger, sorteringsordning, format
  Kolumner           vad varje kolumn betyder och om den är tryckt
  Källa och kontroll varifrån datan kommer och kontrollresultatet
"""
import csv
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROT = Path(__file__).resolve().parent.parent
ARV = ROT / "arv/spelet2_v1"
TRYCKT_TEXT = ROT / "arv/tryckt_text"
KORTDATA = ROT / "kortdata"


def las_csv(relativ_sokvag, nyckel):
    """Läs en arkiverad CSV; rader utan värde i `nyckel` (tomma platshållare) hoppas över."""
    with open(ARV / relativ_sokvag, encoding="utf-8", newline="") as f:
        return [r for r in csv.DictReader(f, delimiter=";") if (r.get(nyckel) or "").strip()]


def tal(v):
    """Tal som tal, allt annat som trimmad text."""
    v = (v or "").strip()
    if re.fullmatch(r"-?\d+", v):
        return int(v)
    if re.fullmatch(r"-?\d+[.,]\d+", v):
        return float(v.replace(",", "."))
    return v


def normalisera(v):
    """För jämförelse CSV ↔ tryck: blanksteg, decimalkomma och tryckets avstavning utjämnas."""
    v = (v or "").replace("­", "").strip()
    v = re.sub(r"(\w) - (\w)", r"\1\2", v)
    v = re.sub(r"\s+", " ", v)
    return v.replace(",", ".")


def _rubrikrad(ws, rubriker):
    ws.append(rubriker)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="4D5E22")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"


def _bredder(ws, standard, specifika=None):
    for i in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(i)].width = standard
    for kol, b in (specifika or {}).items():
        ws.column_dimensions[kol].width = b


def bygg_arbetsbok(*, filnamn, kort_kolumner, rader, namn_kolumn, malltext,
                   produktion_kolumner, format_mm, exemplar, kalla, langa_kolumner=()):
    """Skriv kortdata/<filnamn>.

    kort_kolumner: [(rubrik, csv-kolumn|None, tryckt: bool, förklaring)]
                   Rubrikerna "Kort-id" och "Antal exemplar" fylls automatiskt.
    exemplar:      antal tryckta exemplar per kort (int) eller funktion rad -> int
    kalla:         [(punkt, värde)] till fliken Källa och kontroll
    langa_kolumner: CSV-kolumner med löptext (radbryts, bred kolumn)
    """
    wb = Workbook()

    ws = wb.active
    ws.title = "Kort"
    _rubrikrad(ws, [k[0] for k in kort_kolumner])
    for i, r in enumerate(rader, 1):
        rad = []
        for rubrik, kol, _, _ in kort_kolumner:
            if rubrik == "Kort-id":
                rad.append(i)
            elif rubrik == "Antal exemplar":
                rad.append(exemplar(r) if callable(exemplar) else exemplar)
            elif kol in langa_kolumner:
                rad.append((r.get(kol) or "").strip())
            else:
                rad.append(tal(r.get(kol, "")))
        ws.append(rad)
    specifika = {}
    for idx, (_, kol, _, _) in enumerate(kort_kolumner, 1):
        if kol in langa_kolumner:
            specifika[get_column_letter(idx)] = 60
    _bredder(ws, 14, specifika)
    for rad in ws.iter_rows(min_row=2):
        for c in rad:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    ws = wb.create_sheet("Mallens fasta text")
    _rubrikrad(ws, ["Sida", "Slag", "Text"])
    for m in malltext:
        ws.append(list(m))
    _bredder(ws, 18, {"C": 110})

    ws = wb.create_sheet("Produktion")
    _rubrikrad(ws, ["Kort-id", "Namn"] + produktion_kolumner + ["Format (mm)"])
    for i, r in enumerate(rader, 1):
        ws.append([i, r.get(namn_kolumn, "")] + [tal(r.get(k, "")) for k in produktion_kolumner] + [format_mm])
    _bredder(ws, 14, {"B": 30})

    ws = wb.create_sheet("Kolumner")
    _rubrikrad(ws, ["Kolumn", "Källkolumn i CSV", "Tryckt på kortet", "Förklaring"])
    for rubrik, kol, tryckt, forkl in kort_kolumner:
        ws.append([rubrik, kol or "—", "Ja" if tryckt else "Nej", forkl])
    _bredder(ws, 24, {"D": 70})

    ws = wb.create_sheet("Källa och kontroll")
    _rubrikrad(ws, ["Punkt", "Värde"])
    for p in kalla:
        ws.append(list(p))
    _bredder(ws, 20, {"B": 110})
    for rad in ws.iter_rows(min_row=2):
        rad[1].alignment = Alignment(wrap_text=True, vertical="top")

    ut = KORTDATA / filnamn
    ut.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ut)
    return ut.relative_to(ROT)
