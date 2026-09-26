"""Lägg markexpansionernas former i kortdata/PU_markexpansion.xlsx (kolumnen "Form (rutor)").

Formerna är avlästa ur formbilderna i OneDrive: SPELET 2/Bilder/Former/Markexpansion_NN_Xrutor.png
(NN = Kort-id). En ruta = 250 kvm BYA. Samma format som projektkortens "Form (rutor)": [rad, kolumn].

    python verktyg/lagg_markexpansion_former.py
"""
import json
import sys
from pathlib import Path

import openpyxl

FIL = Path(__file__).resolve().parent.parent / "kortdata" / "PU_markexpansion.xlsx"

FORMER = {
    1: [[0, 0], [1, 0]],
    2: [[0, 0], [1, 0], [2, 0]],
    3: [[0, 0], [0, 1], [1, 1]],
    4: [[0, 0], [0, 1], [1, 1], [1, 2]],
    5: [[0, 0], [0, 1], [1, 0], [2, 0]],
    6: [[0, 0], [0, 1], [1, 0], [1, 1], [2, 0]],
    7: [[0, 0], [0, 1]],
    8: [[0, 0], [1, 0], [1, 1]],
    9: [[0, 0], [0, 1], [1, 1]],
    10: [[0, 0], [0, 1], [0, 2], [1, 0]],
    11: [[0, 0], [0, 1], [1, 0], [1, 1]],
    12: [[0, 0], [0, 1], [1, 0], [1, 1]],
    13: [[0, 0], [1, 0], [2, 0]],
    14: [[0, 0], [0, 1], [1, 0]],
    15: [[0, 0], [0, 1], [0, 2], [1, 0]],
    16: [[0, 0], [0, 1], [1, 0], [1, 1], [2, 1]],
}
KOLUMN = "Form (rutor)"
FORKLARING = ("Nej", "Ej tryckt — markbitens form som [rad, kolumn] per ruta, avläst ur formbilden "
                     "(Bilder/Former/Markexpansion_NN). Antal rutor = BYA / 250")


def main():
    wb = openpyxl.load_workbook(FIL)
    ws = wb["Kort"]
    rubriker = [c.value for c in ws[1]]
    if KOLUMN not in rubriker:
        ws.cell(row=1, column=len(rubriker) + 1, value=KOLUMN)
        rubriker.append(KOLUMN)
    ki, ri, fi = rubriker.index("Kort-id") + 1, rubriker.index("Antal rutor") + 1, rubriker.index(KOLUMN) + 1
    for rad in range(2, ws.max_row + 1):
        kid = ws.cell(row=rad, column=ki).value
        if kid is None:
            continue
        form = FORMER[int(kid)]
        if len(form) != int(ws.cell(row=rad, column=ri).value):
            sys.exit(f"Markexpansion {kid}: formen har {len(form)} rutor, kortet säger {ws.cell(row=rad, column=ri).value}")
        ws.cell(row=rad, column=fi, value=json.dumps(form))
    kol = wb["Kolumner"]
    if KOLUMN not in [c.value for c in kol["A"]]:
        rad = next(r for r in range(2, kol.max_row + 2) if kol.cell(row=r, column=1).value is None)
        kol.cell(row=rad, column=1, value=KOLUMN)
        kol.cell(row=rad, column=2, value="—")
        kol.cell(row=rad, column=3, value=FORKLARING[0])
        kol.cell(row=rad, column=4, value=FORKLARING[1])
    wb.save(FIL)
    print(f"{len(FORMER)} former → {FIL.name}")


if __name__ == "__main__":
    main()
